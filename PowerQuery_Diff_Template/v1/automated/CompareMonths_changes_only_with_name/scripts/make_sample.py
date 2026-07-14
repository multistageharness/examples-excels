"""Generate examples/sample.xlsx -- a workbook shaped the way the diff expects.

Three months as real Excel tables (tbl_Jan, tbl_Feb, tbl_Mar). The data is chosen so
that `xldiff examples/sample.xlsx` reproduces the reference Power Query output:

    Month  ID   Change_Status  Status    Owner
    Feb    103  Removed        Active    Charlie
    Feb    105  Added          Active    Eve
    Mar    102  Removed        Inactive  Bob
    Mar    106  Added          New       Frank

Note Bob (102): Active in Jan, flipped to Inactive in Feb, gone in Mar -- so his Removed
row carries Feb's values, not Jan's. That is the whole point of the "with_name" variant.
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

COLUMNS = ["ID", "Status", "Owner"]

MONTHS = {
    "Jan": [
        [101, "Active", "Alice"],
        [102, "Active", "Bob"],
        [103, "Active", "Charlie"],
        [104, "Active", "Dana"],
    ],
    "Feb": [
        [101, "Active", "Alice"],
        [102, "Inactive", "Bob"],
        [104, "Active", "Dana"],
        [105, "Active", "Eve"],
    ],
    "Mar": [
        [101, "Active", "Alice"],
        [104, "Pending", "Dana"],
        [105, "Active", "Eve"],
        [106, "New", "Frank"],
    ],
}


def build(path: Path) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)

    for month, rows in MONTHS.items():
        worksheet = workbook.create_sheet(month)
        worksheet.append(COLUMNS)
        for row in rows:
            worksheet.append(row)

        table = Table(displayName=f"tbl_{month}", ref=f"A1:C{len(rows) + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        worksheet.add_table(table)

        for column, width in zip("ABC", (8, 12, 14)):
            worksheet.column_dimensions[column].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("examples/sample.xlsx")
    print(f"wrote {build(target)}")
