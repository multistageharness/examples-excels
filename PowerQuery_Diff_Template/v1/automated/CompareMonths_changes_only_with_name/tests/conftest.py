from __future__ import annotations

import sys
from pathlib import Path
from typing import Dict, List, Sequence

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts.make_sample import build as build_sample  # noqa: E402


@pytest.fixture
def sample_workbook(tmp_path: Path) -> Path:
    """The reference workbook: Jan/Feb/Mar as Excel tables, matching the M template."""
    return build_sample(tmp_path / "sample.xlsx")


@pytest.fixture
def make_workbook(tmp_path: Path):
    """Build an .xlsx from {month: rows}, as named tables or as bare sheets."""

    def _make(
        months: Dict[str, List[Sequence]],
        columns: Sequence[str] = ("ID", "Status", "Owner"),
        as_tables: bool = True,
        prefix: str = "tbl_",
        name: str = "book.xlsx",
    ) -> Path:
        workbook = Workbook()
        workbook.remove(workbook.active)

        for month, rows in months.items():
            worksheet = workbook.create_sheet(month)
            worksheet.append(list(columns))
            for row in rows:
                worksheet.append(list(row))

            if as_tables and rows:
                last_column = chr(ord("A") + len(columns) - 1)
                worksheet.add_table(
                    Table(
                        displayName=f"{prefix}{month}",
                        ref=f"A1:{last_column}{len(rows) + 1}",
                    )
                )

        path = tmp_path / name
        workbook.save(path)
        return path

    return _make
